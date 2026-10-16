import io
import shutil
import logging
from pathlib import Path
from app.ingestion.parsers.base import BaseParser, ParsedPage

logger = logging.getLogger(__name__)

OCR_AVAILABLE = shutil.which('tesseract') is not None


class PptxParser(BaseParser):
    def parse(self, file_path: Path):
        from pptx import Presentation
        prs = Presentation(str(file_path))
        pages = []
        for i, slide in enumerate(prs.slides):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                elif shape.shape_type == 13 and OCR_AVAILABLE:  # MSO_SHAPE_TYPE.PICTURE
                    try:
                        from PIL import Image
                        import pytesseract
                        img_bytes = shape.image.blob
                        img = Image.open(io.BytesIO(img_bytes))
                        parts.append(pytesseract.image_to_string(img))
                    except Exception as e:
                        logger.warning('OCR failed for slide %d image: %s', i + 1, e)
            title = ''
            if slide.shapes.title:
                title = slide.shapes.title.text
            if not title:
                title = f'Slide {i + 1}'
            pages.append(ParsedPage(
                page_number=i + 1,
                text='\n'.join(parts),
                metadata={'slide_title': title}
            ))
        return pages
