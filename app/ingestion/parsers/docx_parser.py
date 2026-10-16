import logging
from pathlib import Path
from app.ingestion.parsers.base import BaseParser, ParsedPage

logger = logging.getLogger(__name__)


class DocxParser(BaseParser):
    def parse(self, file_path: Path):
        from docx import Document
        # NOTE: python-docx does not provide a page number API.
        # All DOCX content is assigned page_number=1 because estimating page numbers
        # from heading/paragraph counts would produce false precision.
        # Source citations will show "page 1" for DOCX documents — this is intentional.
        doc = Document(str(file_path))
        parts = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        # Also extract table cell text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        parts.append(text)
        full_text = '\n'.join(parts)
        return [ParsedPage(page_number=1, text=full_text, metadata={})]
