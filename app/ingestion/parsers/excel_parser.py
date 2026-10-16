import logging
from pathlib import Path
from app.ingestion.parsers.base import BaseParser, ParsedPage

logger = logging.getLogger(__name__)


class ExcelParser(BaseParser):
    def parse(self, file_path: Path):
        import openpyxl

        wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
        pages = []

        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            rows_text = []

            for row in ws.iter_rows(values_only=True):
                # Filter out None and empty cells
                cells = [str(c).strip() for c in row if c is not None and str(c).strip() != '']
                if cells:
                    rows_text.append(' | '.join(cells))

            sheet_content = f"Sheet: {sheet_name}\n" + '\n'.join(rows_text) if rows_text else f"Sheet: {sheet_name} (empty)"
            pages.append(ParsedPage(
                page_number=sheet_idx + 1,
                text=sheet_content,
                metadata={'sheet_name': sheet_name}
            ))

        wb.close()
        return pages
